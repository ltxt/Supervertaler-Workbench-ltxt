"""Guard DOCX export against writing translations into the wrong paragraphs.

``import_docx()`` numbers body paragraphs **and** table-cell paragraphs from a
single counter, in document order. ``export_docx()`` used to re-derive that index
with its own counter over ``doc.paragraphs`` (tables skipped), then handle table
cells in a second pass. The two schemes agree only up to the first table: after
it, import's counter has advanced by the number of table-cell paragraphs and
export's has not, so every following body paragraph was overwritten with another
segment's translation. On the DOCX stress-test file, 12 of 62 paragraphs came out
holding the wrong text — and because the surviving text was the untouched
original, the export looked plausible rather than broken.

A second divergence rode along: export tested ``para.text`` for emptiness while
import tested ``_get_full_paragraph_text()``, so a paragraph whose only text sits
inside a hyperlink was imported as a segment but skipped on export.

Both are closed by exporting through the same walk that import numbered by
(``_iter_indexed_paragraphs``). These tests fail loudly if either reappears.
"""

from __future__ import annotations

import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

pytest.importorskip("docx", reason="python-docx is an optional dependency")

from docx import Document  # noqa: E402

from modules.docx_handler import DOCXHandler  # noqa: E402


def _build_docx(path):
    """A document shaped like the ones that broke: paragraphs, then a table,
    then more paragraphs. Anything after the table is where drift shows up."""
    doc = Document()
    doc.add_paragraph("Alpha before the table")
    doc.add_paragraph("Bravo before the table")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cell one"
    table.cell(0, 1).text = "Cell two"
    table.cell(1, 0).text = "Cell three"
    table.cell(1, 1).text = "Cell four"

    doc.add_paragraph("Charlie after the table")
    doc.add_paragraph("Delta after the table")
    doc.add_paragraph("")  # empty: imported as nothing, must not shift indices
    doc.add_paragraph("Echo after the empty paragraph")
    doc.save(str(path))


def _round_trip(tmp_path, transform):
    """Import, apply ``transform(source) -> target`` to every segment, export,
    then re-import. Returns (sources, targets_written, texts_read_back)."""
    src = tmp_path / "source.docx"
    out = tmp_path / "target.docx"
    _build_docx(src)

    handler = DOCXHandler()
    sources = handler.import_docx(str(src), extract_formatting=True)
    segments = [
        {
            "paragraph_id": handler.paragraphs_info[i].paragraph_index,
            "source": text,
            "target": transform(text),
        }
        for i, text in enumerate(sources)
    ]
    handler.export_docx(segments, str(out), preserve_formatting=True)

    reader = DOCXHandler()
    back = reader.import_docx(str(out), extract_formatting=True)
    return sources, [s["target"] for s in segments], back


def test_every_segment_lands_in_its_own_paragraph(tmp_path):
    """Each paragraph must come back holding *its own* translation. Before the
    fix, "Charlie"/"Delta"/"Echo" — everything after the table — held the text of
    an earlier segment instead."""
    sources, targets, back = _round_trip(tmp_path, lambda t: f"[{t}]")

    assert len(back) == len(sources)
    for i, (source, got) in enumerate(zip(sources, back)):
        assert got == f"[{source}]", (
            f"paragraph {i} holds {got!r}, expected the translation of {source!r}"
        )


def test_no_paragraph_is_left_untranslated(tmp_path):
    """The specific symptom: a paragraph silently keeping its source text. A
    marker no reader can mistake makes it unambiguous."""
    _, _, back = _round_trip(tmp_path, lambda t: f"«{t}»")
    untranslated = [i for i, t in enumerate(back) if "«" not in t]
    assert untranslated == [], f"paragraphs still holding source text: {untranslated}"


def test_table_cells_and_body_paragraphs_share_one_index_space(tmp_path):
    """Import's contract, which export has to honour: one counter across body
    paragraphs and table cells, in document order."""
    src = tmp_path / "source.docx"
    _build_docx(src)

    handler = DOCXHandler()
    sources = handler.import_docx(str(src), extract_formatting=True)
    indices = [info.paragraph_index for info in handler.paragraphs_info]

    assert indices == list(range(len(sources))), "import numbering is not dense"
    assert any(info.is_table_cell for info in handler.paragraphs_info)
    # The table sits between "Bravo" and "Charlie", so a correct index space has
    # the cells in the middle rather than appended at the end.
    cell_positions = [i for i, info in enumerate(handler.paragraphs_info)
                      if info.is_table_cell]
    assert cell_positions == [2, 3, 4, 5], cell_positions


def test_export_walk_matches_import_numbering(tmp_path):
    """``_iter_indexed_paragraphs()`` is the shared walk. Its indices and the
    text at each one must agree with what import recorded, or the two halves
    have drifted apart again."""
    src = tmp_path / "source.docx"
    _build_docx(src)

    handler = DOCXHandler()
    handler.import_docx(str(src), extract_formatting=True)

    doc = Document(str(src))
    walked = list(handler._iter_indexed_paragraphs(doc))

    assert [i for i, _ in walked] == list(range(len(handler.paragraphs_info)))
    for index, para in walked:
        expected = handler.paragraphs_info[index].text
        assert handler._get_full_paragraph_text(para).strip() == expected


@pytest.mark.xfail(strict=True, reason="known defect: _replace_paragraph_text "
                                       "does not clear runs nested in <w:hyperlink>")
def test_hyperlink_only_paragraph_is_not_duplicated(tmp_path):
    """A paragraph whose *only* content is a hyperlink comes out doubled.

    ``_replace_paragraph_text()`` writes the translation into the paragraph's
    first direct child run and clears the rest — but a hyperlink's run is nested
    inside ``<w:hyperlink>``, not a direct child, so it is neither reused nor
    cleared. The export ends up holding the source text *and* the translation::

        'Large language model«Large language model»'

    Paragraphs where the hyperlink is only part of the text are unaffected: the
    real corpus file's "Hyperlink field: …" paragraph exports correctly, which is
    why this had not been noticed. Marked ``strict`` so whoever fixes
    ``_replace_paragraph_text`` is told to delete this marker.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    src = tmp_path / "hyperlink.docx"
    out = tmp_path / "hyperlink_target.docx"

    doc = Document()
    doc.add_paragraph("Ordinary paragraph")
    para = doc.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), "rId1")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Large language model"
    run.append(text)
    link.append(run)
    para._p.append(link)
    doc.save(str(src))

    handler = DOCXHandler()
    sources = handler.import_docx(str(src), extract_formatting=True)
    assert "Large language model" in sources, (
        "import no longer sees the hyperlink paragraph; this test needs rewriting")

    segments = [{"paragraph_id": handler.paragraphs_info[i].paragraph_index,
                 "source": t, "target": f"«{t}»"}
                for i, t in enumerate(sources)]
    handler.export_docx(segments, str(out), preserve_formatting=True)

    back = DOCXHandler().import_docx(str(out), extract_formatting=True)
    assert back == ["«Ordinary paragraph»", "«Large language model»"], back
