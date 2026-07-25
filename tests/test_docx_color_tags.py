"""Run colour survives a DOCX round trip as a real inline tag.

Before this, `TagManager` recognised a closed whitelist of tag *names*
(`b|i|u|bi|li|sub|sup`) and had no way to express a formatting property that
carries a **value**. A coloured run therefore produced no tag at all: the
translator never saw the colour and could not move or reproduce it.

Export compensated with a heuristic (`docx_handler._replace_paragraph_text`): it
recorded each original run's colour keyed by that run's *text*, cleared the runs,
rebuilt them from the tagged text, then re-applied colour by matching the text
again. That match only holds while the wording is unchanged — which is to say
never in a real translation. Measured on the project's own DOCX stress-test file:

    original document                  0563C1, C00000
    export with target == source       0563C1, C00000   (a misleading test)
    export with a real translation      C00000 only     -- 0563C1 lost

Colour now travels in the text as ``<cf color="RRGGBB">`` — the same convention
Supervertaler already displays for bilingual CAT formats, and the shape of the
example in the original feature request:

    Specify <cf color="#227acb" font="tahoma"><b>Wall thickness </b></cf>…

so it follows the translator's wording wherever they put it.
"""

import os
import re
import sys
import zipfile

import pytest

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, REPO)

from modules.tag_manager import FormattingRun, TagManager  # noqa: E402


@pytest.fixture
def tm():
    return TagManager()


def _colours(path):
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
    return sorted(set(re.findall(r'<w:color w:val="([0-9A-Fa-f]{6})"', xml)))


# ---------------------------------------------------------------------------
# Emitting the tag
# ---------------------------------------------------------------------------

def test_coloured_run_produces_a_tag_at_all(tm):
    """The defect: a coloured run used to produce no tag, so the colour was
    invisible to the translator and unrecoverable on export."""
    tagged = tm.runs_to_tagged_text([FormattingRun("Wall", color="C00000")])
    assert tagged == '<cf color="C00000">Wall</cf>'


def test_colour_wraps_the_boolean_formatting(tm):
    """Matches the shape in the original request: <cf …><b>text</b></cf>."""
    tagged = tm.runs_to_tagged_text(
        [FormattingRun("Wall thickness", bold=True, color="227ACB")])
    assert tagged == '<cf color="227ACB"><b>Wall thickness</b></cf>'


def test_uncoloured_runs_are_unchanged(tm):
    """Regression guard: existing documents must tag exactly as before."""
    runs = [FormattingRun("Hello "), FormattingRun("world", bold=True),
            FormattingRun("!")]
    assert tm.runs_to_tagged_text(runs) == "Hello <b>world</b>!"


def test_adjacent_runs_of_one_colour_share_a_span(tm):
    tagged = tm.runs_to_tagged_text([
        FormattingRun("Wall ", color="C00000"),
        FormattingRun("thickness", bold=True, color="C00000"),
    ])
    assert tagged == '<cf color="C00000">Wall <b>thickness</b></cf>'


def test_two_colours_produce_two_spans(tm):
    tagged = tm.runs_to_tagged_text([
        FormattingRun("a", color="C00000"),
        FormattingRun("b", color="0563C1"),
    ])
    assert tagged == '<cf color="C00000">a</cf><cf color="0563C1">b</cf>'


def test_colour_span_closes_before_uncoloured_text(tm):
    tagged = tm.runs_to_tagged_text([
        FormattingRun("a", color="C00000"), FormattingRun("b")])
    assert tagged == '<cf color="C00000">a</cf>b'


def test_a_colour_only_run_counts_as_formatted(tm):
    assert FormattingRun("x", color="C00000").has_formatting() is True
    assert FormattingRun("x").has_formatting() is False


# ---------------------------------------------------------------------------
# Reading the tag back
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("runs", [
    [FormattingRun("Wall", color="C00000")],
    [FormattingRun("Wall", bold=True, color="227ACB")],
    [FormattingRun("a", color="C00000"), FormattingRun("b", color="0563C1")],
    [FormattingRun("a", color="C00000"), FormattingRun("b")],
    [FormattingRun("plain")],
    [FormattingRun("Hello "), FormattingRun("w", bold=True, italic=True)],
])
def test_round_trip_preserves_text_and_colour(tm, runs):
    specs = tm.tagged_text_to_runs(tm.runs_to_tagged_text(runs))
    assert "".join(s["text"] for s in specs) == "".join(r.text for r in runs)
    assert ([s.get("color", "") for s in specs if s["text"]]
            == [r.color for r in runs if r.text])


def test_closing_cf_clears_the_colour(tm):
    specs = tm.tagged_text_to_runs('<cf color="C00000">a</cf>b')
    assert [(s["text"], s["color"]) for s in specs] == [("a", "C00000"), ("b", "")]


def test_colour_is_normalised_to_upper_case(tm):
    specs = tm.tagged_text_to_runs('<cf color="c00000">a</cf>')
    assert specs[0]["color"] == "C00000"


def test_a_cf_tag_without_a_colour_is_harmless(tm):
    specs = tm.tagged_text_to_runs("<cf>a</cf>")
    assert [(s["text"], s["color"]) for s in specs] == [("a", "")]


# ---------------------------------------------------------------------------
# The other consumers of TAG_PATTERN
# ---------------------------------------------------------------------------

def test_validate_accepts_a_colour_tag(tm):
    assert tm.validate_tags('<cf color="C00000"><b>x</b></cf>') == (True, "")


def test_validate_still_catches_a_mismatch(tm):
    ok, why = tm.validate_tags('<cf color="C00000"><b>x</cf></b>')
    assert ok is False and why


def test_strip_removes_the_colour_tag(tm):
    assert tm.strip_tags('<cf color="C00000"><b>x</b></cf>') == "x"


def test_count_includes_the_colour_tag(tm):
    counts = tm.count_tags('<cf color="C00000"><b>x</b></cf>')
    assert counts.get("cf") == 1 and counts.get("b") == 1


# ---------------------------------------------------------------------------
# End to end through python-docx
# ---------------------------------------------------------------------------

@pytest.fixture
def coloured_docx(tmp_path):
    docx = pytest.importorskip("docx", reason="python-docx required")
    from docx.shared import RGBColor

    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Specify ")
    run = paragraph.add_run("Wall thickness")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("227ACB")
    paragraph.add_run(" range")
    path = tmp_path / "coloured.docx"
    document.save(str(path))
    return str(path)


def test_import_emits_the_colour_tag(coloured_docx):
    from modules.docx_handler import DOCXHandler

    texts = DOCXHandler().import_docx(coloured_docx, extract_formatting=True)
    assert texts[0] == 'Specify <cf color="227ACB"><b>Wall thickness</b></cf> range'


def test_colour_survives_a_translation_that_shares_no_words(coloured_docx, tmp_path):
    """The case the old heuristic could never handle: every word changes, so
    matching the target text against the source runs cannot re-attach anything."""
    from modules.docx_handler import DOCXHandler

    handler = DOCXHandler()
    texts = handler.import_docx(coloured_docx, extract_formatting=True)
    out = str(tmp_path / "translated.docx")
    handler.export_docx(
        [{"paragraph_id": 0, "source": texts[0],
          "target": 'Geef het <cf color="227ACB"><b>wanddikte</b></cf> bereik'}],
        out)

    assert _colours(out) == ["227ACB"]

    import docx
    runs = [(r.text, r.bold,
             str(r.font.color.rgb) if (r.font.color and r.font.color.rgb) else None)
            for p in docx.Document(out).paragraphs for r in p.runs]
    assert ("wanddikte", True, "227ACB") in runs


def test_without_the_tag_the_colour_is_lost(coloured_docx, tmp_path):
    """Pins the old behaviour, so the value of the tag stays visible: a target
    carrying no <cf …> loses the colour, because there is nothing to carry it."""
    from modules.docx_handler import DOCXHandler

    handler = DOCXHandler()
    texts = handler.import_docx(coloured_docx, extract_formatting=True)
    out = str(tmp_path / "untagged.docx")
    handler.export_docx(
        [{"paragraph_id": 0, "source": texts[0],
          "target": "Geef het <b>wanddikte</b> bereik"}],
        out)
    assert _colours(out) == []


def test_bold_still_survives_alongside_the_colour(coloured_docx, tmp_path):
    from modules.docx_handler import DOCXHandler

    handler = DOCXHandler()
    texts = handler.import_docx(coloured_docx, extract_formatting=True)
    out = str(tmp_path / "bold.docx")
    handler.export_docx(
        [{"paragraph_id": 0, "source": texts[0],
          "target": 'Geef <cf color="227ACB"><b>wanddikte</b></cf> bereik'}],
        out)

    import docx
    bolded = [r.text for p in docx.Document(out).paragraphs
              for r in p.runs if r.bold]
    assert bolded == ["wanddikte"]



@pytest.fixture(scope="session")
def qapp():
    """Session-scoped on purpose.

    A module-scoped QApplication is destroyed when its fixture tears down, which
    takes every QObject with it — including the shared tag renderer — and the
    next test module then cannot register a handler. Every Qt test file in this
    suite therefore keeps the application alive for the whole session.
    """
    pytest.importorskip("PyQt6")
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PyQt6.QtWidgets import QApplication
    yield QApplication.instance() or QApplication([])


@pytest.fixture(scope="module")
def atom_doc(qapp):
    """One document reused across the Qt tests.

    Creating and dropping a QTextDocument per test while a custom text-object
    handler is registered on its layout aborts the interpreter, so it is built
    once here and refilled by each test.
    """
    from PyQt6.QtGui import QTextDocument

    from modules import tag_atoms as ta

    document = QTextDocument()
    ta.register_document(document)
    return document

# ---------------------------------------------------------------------------
# The coupling between the two changesets
# ---------------------------------------------------------------------------
#
# DOCX import (this module) emits the colour tag; the canonical parser and the
# atom layer consume it. They were developed separately and touch no file in
# common, so this is the seam that could silently drift.

def test_a_coloured_docx_run_becomes_a_protected_tag(coloured_docx, atom_doc):
    """The whole chain, in one test: a coloured DOCX run is imported as a real
    attributed tag, parsed by the canonical model, rendered as a protected atom
    at every detail level, and reported by verification when dropped."""
    from PyQt6.QtGui import QTextCursor

    from modules import tag_atoms as ta
    from modules import tag_protection as tp
    from modules.docx_handler import DOCXHandler

    text = DOCXHandler().import_docx(coloured_docx, extract_formatting=True)[0]
    assert '<cf color="227ACB">' in text

    # The canonical parser reads it as an ordinary attributed HTML-family tag,
    # with the colour tag pairing with its own closer.
    tokens = tp.parse_tags(text, number=True)
    assert [(t.name, t.kind, t.number) for t in tokens] == [
        ("cf", tp.KIND_OPEN, 1),
        ("b", tp.KIND_OPEN, 2),
        ("b", tp.KIND_CLOSE, 2),
        ("cf", tp.KIND_CLOSE, 1),
    ]

    document = atom_doc

    # Every detail level renders it and round-trips it byte-exactly.
    for level in ta.DETAIL_LEVELS:
        ta.install_atoms(document, text, detail=level)
        assert ta.document_to_raw(document) == text, level

    # And it is protected: one Backspace removes the entire <cf …> tag rather
    # than leaving a fragment of markup behind.
    ta.install_atoms(document, text, detail=ta.DETAIL_MEDIUM)
    position = ta.atom_tokens(document)[0]["position"]
    cursor = QTextCursor(document)
    cursor.setPosition(position + 1)
    cursor.deletePreviousChar()
    damaged = ta.document_to_raw(document)
    assert "<cf" not in damaged
    assert damaged.count("</cf>") == 1

    issues = tp.verify_tags(text, damaged)
    assert tp.ISSUE_MISSING in {i.kind for i in issues}


def test_the_colour_tag_is_not_mistaken_for_prose(coloured_docx):
    """The genuineness rules must not reject the tag DOCX import just emitted."""
    from modules import tag_protection as tp
    from modules.docx_handler import DOCXHandler

    text = DOCXHandler().import_docx(coloured_docx, extract_formatting=True)[0]
    assert [t.raw for t in tp.parse_tags(text)][0] == '<cf color="227ACB">'
